"""
bgy_exporters/bgy_export_docx.py - Esporta report in formato Word (.docx)
Versione 2.5.0
- colonne, limiti, tema da config_data.json (export)
"""
import os
import pandas as pd
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from datetime import datetime

from bgy_core.bgy_logger import get_logger
from bgy_core.bgy_paths import OUTPUT_DOCX_DIR
from bgy_core.bgy_export_common import (
    get_available_columns, get_max_rows, compute_stats, footer_lines,
)

logger = get_logger("ExportDOCX")
os.makedirs(OUTPUT_DOCX_DIR, exist_ok=True)


def export_to_docx(df, report_type, date_str=None):
    if df.empty:
        logger.warning("⚠️ DataFrame vuoto, nessun file creato")
        return None

    if not date_str:
        date_str = datetime.now().strftime("%Y%m%d")

    max_rows = get_max_rows("docx") or 100

    doc = Document()

    title = doc.add_heading(f"Report {report_type.capitalize()} Voli BGY", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph(
        f"Data: {datetime.now().strftime('%d/%m/%Y %H:%M')}")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(12)

    doc.add_heading("Statistiche", level=2)
    for label, value in compute_stats(df, report_type):
        p = doc.add_paragraph()
        p.add_run(f"{label}: ").bold = True
        p.add_run(str(value))

    doc.add_heading("Dettaglio Voli", level=2)

    cols = get_available_columns(df, report_type)
    if cols:
        table = doc.add_table(rows=1, cols=len(cols))
        table.style = 'Table Grid'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        header_cells = table.rows[0].cells
        for i, col in enumerate(cols):
            header_cells[i].text = col.replace('_', ' ').capitalize()
            if header_cells[i].paragraphs[0].runs:
                header_cells[i].paragraphs[0].runs[0].bold = True

        for idx in range(min(max_rows, len(df))):
            row = table.add_row()
            cells = row.cells
            for i, col in enumerate(cols):
                val = df.iloc[idx][col]
                cells[i].text = str(val) if pd.notna(val) else ""

        if len(df) > max_rows:
            doc.add_paragraph(
                f"... e altri {len(df) - max_rows} voli (non esportati)")

    for line in footer_lines():
        doc.add_paragraph(line)

    filename = f"report_{report_type}_{date_str}.docx"
    filepath = os.path.join(OUTPUT_DOCX_DIR, filename)
    doc.save(filepath)
    logger.info(f"📄 Report Word salvato: {filepath}")
    return filepath


def export_daily_docx(df, date_str=None):
    return export_to_docx(df, "daily", date_str)


def export_nightly_docx(df, date_str=None):
    return export_to_docx(df, "nightly", date_str)


def export_monthly_docx(df, date_str=None):
    return export_to_docx(df, "monthly", date_str)


def export_yearly_docx(df, date_str=None):
    return export_to_docx(df, "yearly", date_str)


if __name__ == "__main__":
    test_df = pd.DataFrame({
        'volo': ['FR123', 'WZ456'],
        'compagnia_aerea': ['Ryanair', 'Wizz Air'],
        'destinazione_origine': ['Londra', 'Parigi'],
    })
    export_to_docx(test_df, "daily", "20260915")