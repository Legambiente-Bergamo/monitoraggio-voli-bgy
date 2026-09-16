"""
bgy_exporters/bgy_export_xlsx.py - Esporta report in formato Excel (.xlsx)
Versione 2.5.0
- colonne, limiti, tema da config_data.json (export)
"""
import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from datetime import datetime

from bgy_core.bgy_logger import get_logger
from bgy_core.bgy_paths import OUTPUT_XLSX_DIR
from bgy_core.bgy_export_common import (
    get_theme, get_available_columns, get_max_rows, compute_stats, footer_lines,
)

logger = get_logger("ExportXLSX")
os.makedirs(OUTPUT_XLSX_DIR, exist_ok=True)


def export_to_xlsx(df, report_type, date_str=None):
    if df.empty:
        logger.warning("⚠️ DataFrame vuoto, nessun file creato")
        return None

    if not date_str:
        date_str = datetime.now().strftime("%Y%m%d")

    theme = get_theme()
    primary = theme.get("primary_color", "1a2a6c")
    text_color = theme.get("text_color", "FFFFFF")

    wb = Workbook()
    ws = wb.active
    ws.title = "Report Voli"

    header_font = Font(bold=True, color=text_color, size=11)
    header_fill = PatternFill(start_color=primary, end_color=primary, fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    cols = get_available_columns(df, report_type)
    export_df = df[cols].copy() if cols else df.copy()

    # Limite righe (0 = nessun limite)
    max_rows = get_max_rows("xlsx")
    if max_rows and max_rows > 0:
        export_df = export_df.head(max_rows)

    for r_idx, row in enumerate(dataframe_to_rows(export_df, index=False, header=True), 1):
        for c_idx, value in enumerate(row, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            cell.border = thin_border
            if r_idx == 1:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if cell.value is not None and len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except Exception:
                pass
        ws.column_dimensions[column_letter].width = min(max_length + 2, 50)

    _add_stats_sheet(wb, df, report_type, theme, max_rows)

    filename = f"report_{report_type}_{date_str}.xlsx"
    filepath = os.path.join(OUTPUT_XLSX_DIR, filename)
    wb.save(filepath)
    logger.info(f"📊 Report Excel salvato: {filepath}")
    return filepath


def _add_stats_sheet(wb, df, report_type, theme, max_rows=0):
    ws = wb.create_sheet("Statistiche")
    primary = theme.get("primary_color", "1a2a6c")
    text_color = theme.get("text_color", "FFFFFF")

    header_font = Font(bold=True, color=text_color, size=11)
    header_fill = PatternFill(start_color=primary, end_color=primary, fill_type="solid")

    ws.cell(row=1, column=1, value="Statistica")
    ws.cell(row=1, column=2, value="Valore")
    for c in (1, 2):
        cell = ws.cell(row=1, column=c)
        cell.font = header_font
        cell.fill = header_fill

    stats = [("Tipo Report", report_type.capitalize())]
    stats.extend(compute_stats(df, report_type))

    if max_rows and max_rows > 0 and len(df) > max_rows:
        stats.append(("", ""))
        stats.append(("Nota", f"Esportate solo {max_rows} righe su {len(df)} totali"))

    stats.append(("", ""))
    for line in footer_lines():
        stats.append(("", line))

    for r_idx, (label, value) in enumerate(stats, 2):
        ws.cell(row=r_idx, column=1, value=label)
        ws.cell(row=r_idx, column=2, value=value)

    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 60


def export_daily_xlsx(df, date_str=None):
    return export_to_xlsx(df, "daily", date_str)


def export_nightly_xlsx(df, date_str=None):
    return export_to_xlsx(df, "nightly", date_str)


def export_monthly_xlsx(df, date_str=None):
    return export_to_xlsx(df, "monthly", date_str)


def export_yearly_xlsx(df, date_str=None):
    return export_to_xlsx(df, "yearly", date_str)


if __name__ == "__main__":
    test_df = pd.DataFrame({
        'volo': ['FR123', 'WZ456'],
        'compagnia_aerea': ['Ryanair', 'Wizz Air'],
        'destinazione_origine': ['Londra', 'Parigi'],
    })
    export_to_xlsx(test_df, "daily", "20260915")