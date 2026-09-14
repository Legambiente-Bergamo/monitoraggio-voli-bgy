"""
exporters/bgy_export_xlsx.py - Esporta report in formato Excel (.xlsx)
"""
import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from datetime import datetime

from core.bgy_logger import get_logger  # <-- CORRETTO
from core.bgy_paths import REPORTS_XLSX_DIR

logger = get_logger("ExportXLSX")
os.makedirs(REPORTS_XLSX_DIR, exist_ok=True)


def export_to_xlsx(df, report_type, date_str=None):
    """
    Esporta un DataFrame in formato Excel con formattazione.
    """
    if df.empty:
        logger.warning("⚠️ DataFrame vuoto, nessun file creato")
        return None
    
    if not date_str:
        date_str = datetime.now().strftime("%Y%m%d")
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Report Voli"
    
    # Stili
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1a2a6c", end_color="1a2a6c", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Scrivi i dati
    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), 1):
        for c_idx, value in enumerate(row, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            cell.border = thin_border
            if r_idx == 1:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
    
    # Auto-adjust columns
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    # Aggiungi foglio statistiche
    add_stats_sheet(wb, df, report_type)
    
    filename = f"report_{report_type}_{date_str}.xlsx"
    filepath = os.path.join(REPORTS_XLSX_DIR, filename)
    wb.save(filepath)
    logger.info(f"📊 Report Excel salvato: {filepath}")
    return filepath


def add_stats_sheet(wb, df, report_type):
    """Aggiunge un foglio con le statistiche"""
    ws_stats = wb.create_sheet("Statistiche")
    
    stats = [
        ("Tipo Report", report_type.capitalize()),
        ("Data Generazione", datetime.now().strftime("%d/%m/%Y %H:%M")),
        ("Totale Voli", len(df)),
        ("", ""),
    ]
    
    if 'tipo_movimento' in df.columns:
        atterraggi = len(df[df['tipo_movimento'].str.contains('Atterraggio|A', na=False)])
        decolli = len(df[df['tipo_movimento'].str.contains('Decollo|D', na=False)])
        stats.append(("Atterraggi", atterraggi))
        stats.append(("Decolli", decolli))
    
    if 'minuti_ritardo' in df.columns:
        ritardi = len(df[df['minuti_ritardo'] > 15])
        in_orario = len(df[df['minuti_ritardo'] <= 15])
        ritardo_medio = df['minuti_ritardo'].mean()
        stats.append(("", ""))
        stats.append(("Voli in Ritardo (>15m)", ritardi))
        stats.append(("Voli in Orario", in_orario))
        stats.append(("Ritardo Medio (min)", f"{ritardo_medio:.1f}"))
    
    for row_idx, (label, value) in enumerate(stats, 1):
        ws_stats.cell(row=row_idx, column=1, value=label)
        ws_stats.cell(row=row_idx, column=2, value=value)
    
    ws_stats.column_dimensions['A'].width = 20
    ws_stats.column_dimensions['B'].width = 30


def export_daily_xlsx(df, date_str=None):
    return export_to_xlsx(df, "daily", date_str)


def export_nightly_xlsx(df, date_str=None):
    return export_to_xlsx(df, "nightly", date_str)


def export_monthly_xlsx(df, date_str=None):
    return export_to_xlsx(df, "monthly", date_str)


if __name__ == "__main__":
    test_df = pd.DataFrame({
        'volo': ['FR123', 'WZ456'],
        'compagnia_aerea': ['Ryanair', 'Wizz Air'],
        'destinazione_origine': ['Londra', 'Parigi']
    })
    export_to_xlsx(test_df, "daily", "20260828")