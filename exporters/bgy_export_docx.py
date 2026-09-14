"""
exporters/bgy_export_docx.py - Esporta report in formato Word (.docx)
"""
import os
import pandas as pd
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from datetime import datetime

from core.bgy_logger import get_logger  # <-- CORRETTO
from core.bgy_paths import REPORTS_DOCX_DIR

logger = get_logger("ExportDOCX")
os.makedirs(REPORTS_DOCX_DIR, exist_ok=True)


def export_to_docx(df, report_type, date_str=None):
    """
    Esporta un DataFrame in formato Word (.docx)
    """
    if df.empty:
        logger.warning("⚠️ DataFrame vuoto, nessun file creato")
        return None
    
    if not date_str:
        date_str = datetime.now().strftime("%Y%m%d")
    
    doc = Document()
    
    # Titolo
    title = doc.add_heading(f"Report {report_type.capitalize()} Voli BGY", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Sottotitolo
    subtitle = doc.add_paragraph(f"Data: {datetime.now().strftime('%d/%m/%Y %H:%M')}")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(12)
    
    # Statistiche
    doc.add_heading("Statistiche", level=2)
    
    stats = [("Totale Voli", len(df))]
    
    if 'tipo_movimento' in df.columns:
        atterraggi = len(df[df['tipo_movimento'].str.contains('Atterraggio|A', na=False)])
        decolli = len(df[df['tipo_movimento'].str.contains('Decollo|D', na=False)])
        stats.append(("Atterraggi", atterraggi))
        stats.append(("Decolli", decolli))
    
    if 'minuti_ritardo' in df.columns:
        ritardi = len(df[df['minuti_ritardo'] > 15])
        in_orario = len(df[df['minuti_ritardo'] <= 15])
        ritardo_medio = df['minuti_ritardo'].mean()
        stats.append(("Voli in Ritardo (>15m)", ritardi))
        stats.append(("Voli in Orario", in_orario))
        stats.append(("Ritardo Medio (min)", f"{ritardo_medio:.1f}"))
    
    for label, value in stats:
        p = doc.add_paragraph()
        p.add_run(f"{label}: ").bold = True
        p.add_run(str(value))
    
    # Tabella dei voli
    doc.add_heading("Dettaglio Voli", level=2)
    
    cols_to_show = ['volo', 'callsign', 'tipo_movimento', 'compagnia_aerea', 
                    'destinazione_origine', 'orario_schedulato', 'orario_effettivo', 
                    'minuti_ritardo', 'stato_ritardo']
    cols = [c for c in cols_to_show if c in df.columns]
    
    table = doc.add_table(rows=1, cols=len(cols))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    header_cells = table.rows[0].cells
    for i, col in enumerate(cols):
        header_cells[i].text = col.replace('_', ' ').capitalize()
        header_cells[i].paragraphs[0].runs[0].bold = True
    
    max_rows = min(100, len(df))
    for idx in range(max_rows):
        row = table.add_row()
        cells = row.cells
        for i, col in enumerate(cols):
            cells[i].text = str(df.iloc[idx][col]) if pd.notna(df.iloc[idx][col]) else ""
    
    if len(df) > 100:
        doc.add_paragraph(f"... e altri {len(df) - 100} voli")
    
    # Footer
    doc.add_paragraph(f"Report generato automaticamente il {datetime.now().strftime('%d/%m/%Y %H:%M')}")
    doc.add_paragraph("Dati elaborati da fonti pubbliche (SACBO e OpenSky Network)")
    
    filename = f"report_{report_type}_{date_str}.docx"
    filepath = os.path.join(REPORTS_DOCX_DIR, filename)
    doc.save(filepath)
    logger.info(f"📄 Report Word salvato: {filepath}")
    return filepath


def export_daily_docx(df, date_str=None):
    return export_to_docx(df, "daily", date_str)


def export_nightly_docx(df, date_str=None):
    return export_to_docx(df, "nightly", date_str)


def export_monthly_docx(df, date_str=None):
    return export_to_docx(df, "monthly", date_str)


if __name__ == "__main__":
    test_df = pd.DataFrame({
        'volo': ['FR123', 'WZ456'],
        'compagnia_aerea': ['Ryanair', 'Wizz Air'],
        'destinazione_origine': ['Londra', 'Parigi']
    })
    export_to_docx(test_df, "daily", "20260828")